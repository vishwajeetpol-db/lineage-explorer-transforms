"""Unit tests for backend/excel_export.py.

Exercises the public build_lineage_workbook entry point plus the pure layout
helpers, covering: full multi-sheet build (schema scope with a single Lineage
Map), catalog scope (per-schema maps), empty data, orphan handling, cycle
removal in the layering, and the entity-edge collapse helper. Everything is
in-memory (openpyxl only) — no network.
"""
import io

import pytest
from openpyxl import load_workbook

from backend import excel_export as ee
from backend.models import (
    ColumnLineageEdge,
    ColumnLineageResponse,
    EntityNode,
    LineageEdge,
    LineageResponse,
    TableNode,
)


# --------------------------------------------------------------------------
# helpers to build fixtures
# --------------------------------------------------------------------------
def _table(name, schema="sales", status="connected", **kw):
    fq = f"cat.{schema}.{name}"
    return TableNode(
        id=fq,
        name=name,
        full_name=fq,
        table_type=kw.pop("table_type", "MANAGED"),
        owner=kw.pop("owner", "alice@x.com"),
        comment=kw.pop("comment", "a table"),
        columns=kw.pop("columns", [{"name": "c1"}, {"name": "c2"}]),
        created_at=kw.pop("created_at", "2024-01-01"),
        updated_at=kw.pop("updated_at", "2024-02-01"),
        upstream_count=kw.pop("upstream_count", 1),
        downstream_count=kw.pop("downstream_count", 1),
        lineage_status=status,
    )


def _entity(etype="JOB", eid="123456789", **kw):
    return EntityNode(
        id=f"entity:{etype}:{eid}",
        entity_type=etype,
        entity_id=eid,
        display_name=kw.pop("display_name", "nightly-etl"),
        last_run=kw.pop("last_run", "2024-02-01T00:00:00Z"),
        owner=kw.pop("owner", "bob@x.com"),
        cost_usd=kw.pop("cost_usd", 12.5),
    )


def _load(data: bytes):
    return load_workbook(io.BytesIO(data))


# --------------------------------------------------------------------------
# pure helpers
# --------------------------------------------------------------------------
def test_is_entity():
    assert ee._is_entity("entity:JOB:1")
    assert not ee._is_entity("cat.s.t")


def test_split_fqdn():
    assert ee._split_fqdn("cat.sales.orders") == ("cat", "sales")
    assert ee._split_fqdn("weird") == ("", "")


def test_schema_of():
    assert ee._schema_of("cat.sales.orders") == "sales"
    assert ee._schema_of("bad") == ""


def test_collapse_edges_direct_and_via_entity():
    class E:
        def __init__(self, s, t):
            self.source, self.target = s, t

    edges = [
        E("cat.s.a", "cat.s.b"),               # direct table->table
        E("cat.s.a", "entity:JOB:1"),          # table -> entity
        E("entity:JOB:1", "cat.s.c"),          # entity -> table
        E("cat.s.a", "cat.s.a"),               # self loop, dropped
    ]
    out = ee._collapse_edges(edges)
    assert ("cat.s.a", "cat.s.b") in out
    assert ("cat.s.a", "cat.s.c") in out       # via entity JOB:1
    assert all(s != t for s, t in out)


def test_layer_nodes_with_cycle_and_orphan():
    a = _table("a")
    b = _table("b")
    c = _table("c")
    orphan = _table("orphan_tbl")
    ids = {n.id: n for n in [a, b, c, orphan]}
    # a -> b -> c -> a  (cycle) ; orphan has no edges
    edges = [(a.id, b.id), (b.id, c.id), (c.id, a.id)]
    layers, orphans, adj = ee._layer_nodes(ids, edges)
    assert orphan.id in orphans
    assert a.id not in orphans
    # every connected node placed in some layer
    placed = {n for nodes in layers.values() for n in nodes}
    assert {a.id, b.id, c.id} <= placed


# --------------------------------------------------------------------------
# build_lineage_workbook — full schema-scoped build
# --------------------------------------------------------------------------
def test_build_schema_scope_full():
    t1 = _table("orders", status="root")
    t2 = _table("orders_agg", status="connected")
    t3 = _table("orders_final", status="leaf")
    ent = _entity()
    nodes = [t1, t2, t3, ent]
    edges = [
        LineageEdge(source=t1.id, target=ent.id),
        LineageEdge(source=ent.id, target=t2.id),
        LineageEdge(source=t2.id, target=t3.id),
    ]
    resp = LineageResponse(nodes=nodes, edges=edges)
    col_edges = ColumnLineageResponse(edges=[
        ColumnLineageEdge(source_table=t1.id, source_column="c1",
                          target_table=t2.id, target_column="c1"),
    ])
    table_edges = [
        {"source": t1.id, "target": t2.id, "entity_type": "JOB", "entity_id": "123456789"},
        {"source": t2.id, "target": t3.id, "entity_type": None, "entity_id": None},
    ]
    data = ee.build_lineage_workbook(
        "cat", "sales", resp, column_edges=col_edges,
        entity_names={ent.id: "Nightly ETL"}, table_edges=table_edges,
    )
    assert isinstance(data, bytes) and len(data) > 0
    wb = _load(data)
    assert "Summary" in wb.sheetnames
    assert "Tables" in wb.sheetnames
    assert "Lineage" in wb.sheetnames
    assert "Pipelines" in wb.sheetnames
    assert "Column Lineage" in wb.sheetnames
    assert "Lineage Map" in wb.sheetnames
    # Tables sheet has a header + 3 data rows
    tbl = wb["Tables"]
    assert tbl.max_row == 4


def test_build_catalog_scope_per_schema_maps():
    # two schemas so per-schema map path runs
    a = _table("a", schema="sales", status="root")
    b = _table("b", schema="sales", status="leaf")
    c = _table("c", schema="mkt", status="root")
    d = _table("d", schema="mkt", status="leaf")
    ent = _entity()
    nodes = [a, b, c, d, ent]
    edges = [
        LineageEdge(source=a.id, target=b.id),
        LineageEdge(source=c.id, target=d.id),
    ]
    resp = LineageResponse(nodes=nodes, edges=edges)
    table_edges = [
        {"source": a.id, "target": b.id, "entity_type": "JOB", "entity_id": "999"},
        {"source": c.id, "target": d.id, "entity_type": "JOB", "entity_id": "999"},
    ]
    # schema=None => catalog scope => _add_per_schema_maps
    data = ee.build_lineage_workbook("cat", None, resp, table_edges=table_edges,
                                     entity_names={ent.id: "j"})
    wb = _load(data)
    map_sheets = [s for s in wb.sheetnames if s.startswith("Map")]
    assert len(map_sheets) >= 1


def test_build_empty_data():
    resp = LineageResponse(nodes=[], edges=[])
    data = ee.build_lineage_workbook("cat", "sales", resp)
    assert isinstance(data, bytes) and len(data) > 0
    wb = _load(data)
    # Summary always present; no entity/column sheets when empty
    assert "Summary" in wb.sheetnames
    assert "Tables" in wb.sheetnames
    assert "Pipelines" not in wb.sheetnames
    assert "Column Lineage" not in wb.sheetnames


def test_build_with_orphans_and_no_entity_names():
    t1 = _table("connected_a", status="root")
    t2 = _table("connected_b", status="leaf")
    orphan = _table("lonely", status="orphan")
    ent = _entity(display_name=None)  # forces fallback naming
    nodes = [t1, t2, orphan, ent]
    edges = [LineageEdge(source=t1.id, target=t2.id)]
    resp = LineageResponse(nodes=nodes, edges=edges)
    table_edges = [{"source": t1.id, "target": t2.id, "entity_type": "JOB", "entity_id": "abcdef12345"}]
    data = ee.build_lineage_workbook("cat", "sales", resp, table_edges=table_edges)
    wb = _load(data)
    assert "Lineage Map" in wb.sheetnames
    # Pipelines sheet exists because there is an entity node
    assert "Pipelines" in wb.sheetnames


def test_build_lineage_map_sheet_empty_nodes_returns_early():
    from openpyxl import Workbook
    wb = Workbook()
    before = list(wb.sheetnames)
    ee._build_lineage_map_sheet(wb, [], [], {}, "Lineage Map")
    assert wb.sheetnames == before  # no sheet added for empty node set


def test_entity_via_label_and_cost_none():
    ent = _entity(cost_usd=None, display_name="J")
    resp = LineageResponse(nodes=[ent], edges=[])
    data = ee.build_lineage_workbook("cat", "sales", resp)
    wb = _load(data)
    pipe = wb["Pipelines"]
    # header + 1 entity row
    assert pipe.max_row == 2
